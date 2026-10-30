"""rsi_report.py — يبني القائمة والتنبيهات (xlsx) من الحالة. كل أسهم الحالة تظهر في القائمة بلا استبعاد."""
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from rsi_engine import zone, classify, YEARS

def rows(state):
    out = []
    for t, s in state.items():
        yl, yh = zone(s); x = s.get('rsi') if 'pending' not in s else None; red = s.get('hist_low_rsi')
        ok = x is not None
        out.append({'الرمز': t, 'RSI_الحالي': round(x, 2) if ok else None,
          'السعر_الحالي': round(s['last_close'], 4 if s['last_close'] < 1 else 2),
          'القاع_الاحمر': round(red, 2) if red is not None else None, 'تاريخ_القاع_الاحمر': s.get('hist_low_date'),
          'الاصفر_ادنى': round(yl, 2) if yl is not None else None, 'الاصفر_اعلى': round(yh, 2) if yh is not None else None,
          'قريب_من_المنطقة_الآن': ('نعم' if (yh is not None and x <= yh + 1e-9) else 'لا') if ok else None,
          'نوع_المنطقة_الحالية': classify(x, red, yl, yh) if ok else '-',
          **{f'قاع_RSI_{y}': (round(s['year_lows'][y], 2) if ok and y in s.get('year_lows', {}) else None) for y in YEARS[:7]},
          'متوسط_حجم_التداول': int(s.get('median_volume', 0))})
    df = pd.DataFrame(out)
    df['_k'] = df['نوع_المنطقة_الحالية'].map(lambda z: 0 if z.startswith('أحمر') else (1 if z == 'أصفر' else 2))
    return df.sort_values(['_k', 'RSI_الحالي', 'الرمز'], na_position='last').drop(columns='_k').reset_index(drop=True)

FILLS = {k: 'FFE699' for k in ['الاصفر_ادنى', 'الاصفر_اعلى', 'الاصفر_الادنى', 'الاصفر_الاعلى']}
FILLS.update({'القاع_الاحمر': 'F4B6B6', 'تاريخ_القاع_الاحمر': 'F4B6B6'})

def _sheet(wb, name, d, widths):
    F = lambda **k: Font(name='Arial', size=10, **k)
    w = wb.create_sheet(name); w.sheet_view.rightToLeft = True
    for j, col in enumerate(d.columns, 1):
        c = w.cell(1, j, col); c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        c.fill = PatternFill('solid', fgColor=FILLS.get(col, '1F3864')); c.font = F(bold=True, color='000000' if col in FILLS else 'FFFFFF')
        w.column_dimensions[get_column_letter(j)].width = widths.get(col, 14)
    for i, row in enumerate(d.itertuples(index=False), 2):
        for j, v in enumerate(row, 1):
            if pd.isna(v): v = None
            c = w.cell(i, j, v); c.font = F(); c.alignment = Alignment(horizontal='center')
    w.row_dimensions[1].height = 32; w.freeze_panes = 'B2'
    w.auto_filter.ref = f'A1:{get_column_letter(max(1, len(d.columns)))}{len(d) + 1}'

def write_alerts_xlsx(path, results):
    """ملف التنبيهات اليومي: ورقة لكل محرك. results = [(NAME, TIER, DataFrame, error_or_None)].
    إذا فشل محرك تظهر ورقته برسالة الخطأ (فيُعرف أي محرك يُصلَح). يرجع {TIER: عدد}."""
    wb = Workbook(); wb.remove(wb.active)
    for name, tier, df, err in results:
        if err:
            w = wb.create_sheet(name); w.sheet_view.rightToLeft = True
            w.cell(1, 1, f'❌ فشل هذا المحرك: {err}').font = Font(name='Arial', size=11, bold=True, color='C00000')
            w.column_dimensions['A'].width = 120
        else:
            _sheet(wb, name, df, {'نوع_التنبيه': 44, 'تنبيه_سيولة': 26, 'ملاحظة': 36, 'آخر_يوم_تداول': 16})
    wb.save(path)
    return {tier: (None if err else len(df)) for _, tier, df, err in results}

def write_list_xlsx(path, df):
    """القائمة المركزية (كل السوق) — تُرسل شهرياً."""
    wb = Workbook(); wb.remove(wb.active)
    _sheet(wb, '1_القائمة_الرئيسية', df, {'نوع_المنطقة_الحالية': 24, 'تاريخ_القاع_الاحمر': 14, 'متوسط_حجم_التداول': 16})
    wb.save(path)
